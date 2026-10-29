"""
Parser e validador offline de pré-visualização de importação CSV/XLSX.

Funções puras: recebem bytes já carregados em memória e devolvem uma
pré-visualização normalizada. Não acessam banco, rede, disco ou logs.
Nenhum registro financeiro é criado aqui; a confirmação é outro contrato.

Mensagens de erro usam apenas códigos fixos, número de linha e nome do
campo canônico, sem ecoar conteúdo do arquivo (dados pessoais).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date
from datetime import datetime
from decimal import Decimal
from decimal import InvalidOperation
from typing import Any

import openpyxl

from app.core.money import MAX_MONEY
from app.core.money import parse_localized_money


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_COLUMNS = 30
MAX_SHEETS = 1
MAX_CELL_CHARS = 1000

MAX_ZIP_ENTRIES = 200
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100

CLIENTE_MIN_CHARS = 2
CLIENTE_MAX_CHARS = 150
EMAIL_MAX_CHARS = 150
WHATSAPP_MAX_CHARS = 30
WHATSAPP_MIN_DIGITS = 8

REQUIRED_FIELDS = ("cliente", "valor", "vencimento", "status")
OPTIONAL_FIELDS = ("email", "whatsapp")
CANONICAL_FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS

HEADER_ALIASES = {
    "cliente": "cliente",
    "nome": "cliente",
    "nome do cliente": "cliente",
    "valor": "valor",
    "vencimento": "vencimento",
    "data de vencimento": "vencimento",
    "data vencimento": "vencimento",
    "status": "status",
    "situacao": "status",
    "email": "email",
    "e mail": "email",
    "whatsapp": "whatsapp",
}

CSV_DELIMITERS = (";", ",")

FILE_ERROR_MESSAGES = {
    "empty_file": "O arquivo está vazio.",
    "file_too_large": "O arquivo excede o tamanho máximo permitido.",
    "unsupported_extension": "Extensão não suportada. Use .csv ou .xlsx.",
    "content_mismatch": "O conteúdo não corresponde à extensão do arquivo.",
    "unsupported_encoding": "Codificação não suportada. Salve o CSV em UTF-8.",
    "binary_content": "O arquivo contém caracteres binários não permitidos.",
    "malformed_csv": "O CSV está malformado.",
    "ambiguous_delimiter": "Não foi possível determinar o delimitador do CSV.",
    "missing_header": "O arquivo não possui cabeçalho.",
    "blank_header": "Há coluna sem nome no cabeçalho.",
    "missing_required_fields": "Campos obrigatórios ausentes no cabeçalho.",
    "ambiguous_header": "Mais de uma coluna corresponde ao mesmo campo.",
    "too_many_columns": "O arquivo excede o número máximo de colunas.",
    "too_many_rows": "O arquivo excede o número máximo de linhas.",
    "cell_too_long": "Uma célula excede o tamanho máximo permitido.",
    "no_data_rows": "O arquivo não possui linhas de dados.",
    "malformed_zip": "O arquivo XLSX está corrompido ou não é um ZIP válido.",
    "unsafe_zip": "O arquivo XLSX excede os limites de segurança.",
    "encrypted_workbook": "Planilhas protegidas por senha não são suportadas.",
    "macro_workbook": "Planilhas com macros não são suportadas.",
    "external_links": "Planilhas com vínculos externos não são suportadas.",
    "formula_workbook": "Planilhas com fórmulas não são suportadas.",
    "unsafe_xml": "O arquivo XLSX contém XML não permitido.",
    "too_many_sheets": "A planilha deve conter exatamente uma aba.",
    "malformed_workbook": "Não foi possível ler a planilha XLSX.",
    "invalid_mapping": "O mapeamento manual de colunas é inválido.",
    "invalid_delimiter": "Delimitador do CSV ausente, inválido ou não permitido.",
}

ROW_ISSUE_MESSAGES = {
    "required": "Campo obrigatório vazio.",
    "invalid_type": "Tipo de valor não suportado para o campo.",
    "too_short": "Valor menor que o tamanho mínimo.",
    "too_long": "Valor maior que o tamanho máximo.",
    "invalid_characters": "Valor contém caracteres não permitidos.",
    "invalid_email": "E-mail inválido.",
    "invalid_whatsapp": "WhatsApp inválido.",
    "invalid_money": "Valor monetário inválido.",
    "ambiguous_money": "Valor monetário ambíguo; use separador decimal explícito.",
    "non_positive_money": "O valor deve ser maior que zero.",
    "money_out_of_range": "O valor excede o limite suportado.",
    "invalid_date": "Data inválida. Use AAAA-MM-DD ou DD/MM/AAAA.",
    "status_not_allowed": "Status inicial não permitido; use aberto.",
    "invalid_status": "Status desconhecido; use aberto.",
    "extra_columns": "A linha possui mais colunas que o cabeçalho.",
    "duplicate_in_file": "Linha duplicada no arquivo (cliente, valor e vencimento).",
}

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_CONTROL_CHARS = re.compile(r"[\x00-\x1f\x7f]")
_EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_WHATSAPP_PATTERN = re.compile(r"^\+?[0-9(][0-9 ()\-]*$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_BR_DATE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
_MONEY_PLAIN = re.compile(r"^\d+(?:[.,]\d{1,2})?$")
_MONEY_GROUPED_BR = re.compile(r"^\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?$")
_MONEY_GROUPED_US = re.compile(r"^\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?$")
_SHEET_ENTRY = re.compile(r"^xl/worksheets/[^/]+\.xml$")
_FORMULA_TAG = re.compile(rb"<(?:[A-Za-z_][\w.\-]*:)?f[\s>/]")
_UNSAFE_XML = re.compile(rb"<!(?:DOCTYPE|ENTITY)", re.IGNORECASE)
_HEADER_SEPARATORS = re.compile(r"[^a-z0-9]+")

_XLSX_SHEET_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.worksheet+xml"
)


class ImportPreviewError(ValueError):
    """
    Rejeição do arquivo inteiro. Carrega apenas um código fixo.
    """

    def __init__(self, code: str) -> None:
        if code not in FILE_ERROR_MESSAGES:
            raise ValueError("Código de erro de importação desconhecido.")

        self.code = code
        self.message = FILE_ERROR_MESSAGES[code]

        super().__init__(self.message)

    def __repr__(self) -> str:
        return f"ImportPreviewError(code={self.code!r})"


@dataclass(frozen=True)
class RowIssue:
    field: str | None
    code: str

    @property
    def message(self) -> str:
        return ROW_ISSUE_MESSAGES[self.code]


@dataclass(frozen=True)
class PreviewRow:
    row_number: int
    values: dict[str, Any] | None
    issues: tuple[RowIssue, ...]

    @property
    def is_valid(self) -> bool:
        return not self.issues


@dataclass(frozen=True)
class ImportPreview:
    file_kind: str
    file_sha256: str
    content_sha256: str
    column_mapping: dict[str, int]
    ignored_columns: tuple[int, ...]
    rows: tuple[PreviewRow, ...]

    @property
    def total_rows(self) -> int:
        return len(self.rows)

    @property
    def valid_rows(self) -> int:
        return sum(1 for row in self.rows if row.is_valid)

    @property
    def invalid_rows(self) -> int:
        return self.total_rows - self.valid_rows

    @property
    def valid_total(self) -> Decimal:
        return sum(
            (row.values["valor"] for row in self.rows if row.is_valid),
            Decimal("0.00"),
        )


def parse_import_preview(
    filename: str,
    data: bytes,
    *,
    column_mapping: dict[str, int] | None = None,
    csv_delimiter: str | None = None,
) -> ImportPreview:
    """
    Ponto de entrada único. Não persiste nada e não registra logs.

    Sem ``column_mapping``, o cabeçalho é reconhecido por aliases. Com
    ``column_mapping`` (campo canônico -> índice zero-based), o mapeamento
    manual substitui os aliases; CSV exige ``csv_delimiter`` explícito e
    XLSX o proíbe.
    """

    if not isinstance(data, (bytes, bytearray)):
        raise TypeError("data deve ser bytes.")

    data = bytes(data)

    if not data:
        raise ImportPreviewError("empty_file")

    if len(data) > MAX_FILE_BYTES:
        raise ImportPreviewError("file_too_large")

    kind = _detect_kind(filename)
    manual = _check_manual_options(kind, column_mapping, csv_delimiter)

    if kind == "csv":
        table = _read_csv_table(data, csv_delimiter)
    else:
        table = _read_xlsx_table(data)

    if manual is None:
        mapping, ignored = _map_header(table[0])
    else:
        mapping, ignored = _apply_manual_mapping(table[0], manual)
    rows = _validate_rows(table[1:], mapping, header_width=len(table[0]))

    if not rows:
        raise ImportPreviewError("no_data_rows")

    return ImportPreview(
        file_kind=kind,
        file_sha256=hashlib.sha256(data).hexdigest(),
        content_sha256=_content_digest(mapping, rows),
        column_mapping=mapping,
        ignored_columns=ignored,
        rows=rows,
    )


def _detect_kind(filename: str) -> str:
    name = str(filename or "").strip().lower()

    if name.endswith(".csv"):
        return "csv"

    if name.endswith(".xlsx"):
        return "xlsx"

    raise ImportPreviewError("unsupported_extension")


def _check_manual_options(
    kind: str,
    column_mapping: Any,
    csv_delimiter: Any,
) -> dict[str, int] | None:
    """
    Valida a forma do mapeamento manual e do delimitador antes de ler o
    arquivo. Devolve uma cópia em ordem canônica, ou None sem mapeamento.
    """

    if column_mapping is None:
        if csv_delimiter is not None:
            raise ImportPreviewError("invalid_delimiter")

        return None

    if kind == "xlsx":
        if csv_delimiter is not None:
            raise ImportPreviewError("invalid_delimiter")
    elif not isinstance(csv_delimiter, str) or csv_delimiter not in CSV_DELIMITERS:
        raise ImportPreviewError("invalid_delimiter")

    if not isinstance(column_mapping, dict):
        raise ImportPreviewError("invalid_mapping")

    fields = set(column_mapping)

    if not set(REQUIRED_FIELDS) <= fields <= set(CANONICAL_FIELDS):
        raise ImportPreviewError("invalid_mapping")

    indices = list(column_mapping.values())

    # type() exclui bool, que é subclasse de int.
    if any(type(index) is not int or index < 0 for index in indices):
        raise ImportPreviewError("invalid_mapping")

    if len(set(indices)) != len(indices):
        raise ImportPreviewError("invalid_mapping")

    return {
        field: column_mapping[field]
        for field in CANONICAL_FIELDS
        if field in column_mapping
    }


# ---------------------------------------------------------------- CSV


def _read_csv_table(data: bytes, delimiter: str | None = None) -> list[list[Any]]:
    if data.startswith(b"PK\x03\x04"):
        raise ImportPreviewError("content_mismatch")

    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]

    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as error:
        raise ImportPreviewError("unsupported_encoding") from error

    if "\x00" in text:
        raise ImportPreviewError("binary_content")

    if delimiter is not None:
        table = _parse_csv_text(text, delimiter)

        if table is None:
            raise ImportPreviewError("malformed_csv")

        return _bound_table(table)

    candidates = []

    for delimiter in CSV_DELIMITERS:
        table = _parse_csv_text(text, delimiter)

        if table is None or not table:
            continue

        header = _trim_trailing_blanks(table[0])
        recognized = {
            HEADER_ALIASES.get(_normalize_header(cell))
            for cell in header
        }

        if set(REQUIRED_FIELDS) <= recognized:
            candidates.append(table)

    if len(candidates) > 1:
        raise ImportPreviewError("ambiguous_delimiter")

    if not candidates:
        table = _parse_csv_text(text, CSV_DELIMITERS[0])

        if table is None:
            raise ImportPreviewError("malformed_csv")

        if not table:
            raise ImportPreviewError("missing_header")

        raise ImportPreviewError("missing_required_fields")

    return _bound_table(candidates[0])


def _parse_csv_text(text: str, delimiter: str) -> list[list[str]] | None:
    reader = csv.reader(
        io.StringIO(text, newline=""),
        delimiter=delimiter,
        strict=True,
    )
    table = []

    try:
        for record in reader:
            if _is_blank_record(record):
                if table:
                    table.append(None)
                continue

            table.append(record)

            if len(table) > MAX_DATA_ROWS + 1:
                raise ImportPreviewError("too_many_rows")

    except csv.Error:
        return None

    return table


# ---------------------------------------------------------------- XLSX


def _read_xlsx_table(data: bytes) -> list[list[Any]]:
    if not data.startswith(b"PK\x03\x04"):
        raise ImportPreviewError("content_mismatch")

    _inspect_xlsx_zip(data)

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except ImportPreviewError:
        raise
    except Exception as error:
        raise ImportPreviewError("malformed_workbook") from error

    try:
        if len(workbook.sheetnames) != MAX_SHEETS:
            raise ImportPreviewError("too_many_sheets")

        sheet = workbook.worksheets[0]
        table = []

        for record in sheet.iter_rows(values_only=True):
            record = list(record)

            if _is_blank_record(record):
                if table:
                    table.append(None)
                continue

            table.append(record)

            if len(table) > MAX_DATA_ROWS + 1:
                raise ImportPreviewError("too_many_rows")

    except ImportPreviewError:
        raise
    except Exception as error:
        raise ImportPreviewError("malformed_workbook") from error
    finally:
        workbook.close()

    if not table:
        raise ImportPreviewError("missing_header")

    return _bound_table(table)


def _inspect_xlsx_zip(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError) as error:
        raise ImportPreviewError("malformed_zip") from error

    with archive:
        try:
            entries = archive.infolist()
        except Exception as error:
            raise ImportPreviewError("malformed_zip") from error

        if len(entries) > MAX_ZIP_ENTRIES:
            raise ImportPreviewError("unsafe_zip")

        names = set()
        total = 0

        for entry in entries:
            name = entry.filename

            if name in names:
                raise ImportPreviewError("unsafe_zip")

            names.add(name)

            if entry.flag_bits & 0x1:
                raise ImportPreviewError("encrypted_workbook")

            total += entry.file_size

            if total > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ImportPreviewError("unsafe_zip")

            if entry.file_size > MAX_ZIP_COMPRESSION_RATIO * max(
                entry.compress_size, 1
            ) and entry.file_size > 1024 * 1024:
                raise ImportPreviewError("unsafe_zip")

            lowered = name.lower()

            if lowered.endswith("vbaproject.bin") or "macrosheets/" in lowered:
                raise ImportPreviewError("macro_workbook")

            if lowered.startswith("xl/externallinks/"):
                raise ImportPreviewError("external_links")

        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ImportPreviewError("malformed_workbook")

        for entry in entries:
            lowered = entry.filename.lower()

            if not (lowered.endswith(".xml") or lowered.endswith(".rels")):
                continue

            content = _read_zip_entry(archive, entry)

            if _UNSAFE_XML.search(content):
                raise ImportPreviewError("unsafe_xml")

            if entry.filename == "[Content_Types].xml":
                if b"macroEnabled" in content:
                    raise ImportPreviewError("macro_workbook")

            if _SHEET_ENTRY.match(entry.filename) and _FORMULA_TAG.search(content):
                raise ImportPreviewError("formula_workbook")

        sheet_entries = [n for n in names if _SHEET_ENTRY.match(n)]

        if len(sheet_entries) > MAX_SHEETS:
            raise ImportPreviewError("too_many_sheets")

        if any(n.lower().startswith("xl/chartsheets/") for n in names):
            raise ImportPreviewError("too_many_sheets")


def _read_zip_entry(archive: zipfile.ZipFile, entry: zipfile.ZipInfo) -> bytes:
    limit = MAX_ZIP_UNCOMPRESSED_BYTES

    try:
        with archive.open(entry) as handle:
            content = handle.read(limit + 1)
    except Exception as error:
        raise ImportPreviewError("malformed_zip") from error

    if len(content) > limit or len(content) != entry.file_size:
        raise ImportPreviewError("unsafe_zip")

    return content


# ---------------------------------------------------------------- tabela


def _is_blank_record(record: list[Any]) -> bool:
    return all(_is_blank_cell(cell) for cell in record)


def _is_blank_cell(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _trim_trailing_blanks(record: list[Any]) -> list[Any]:
    end = len(record)

    while end and _is_blank_cell(record[end - 1]):
        end -= 1

    return list(record[:end])


def _bound_table(table: list[list[Any] | None]) -> list[list[Any] | None]:
    while table and table[-1] is None:
        table.pop()

    if not table or table[0] is None:
        raise ImportPreviewError("missing_header")

    if len(table) - 1 > MAX_DATA_ROWS:
        raise ImportPreviewError("too_many_rows")

    for record in table:
        if record is None:
            continue

        if len(_trim_trailing_blanks(record)) > MAX_COLUMNS:
            raise ImportPreviewError("too_many_columns")

        for cell in record:
            if isinstance(cell, str) and len(cell) > MAX_CELL_CHARS:
                raise ImportPreviewError("cell_too_long")

    table[0] = _trim_trailing_blanks(table[0])

    return table


def _normalize_header(value: Any) -> str:
    if not isinstance(value, str):
        return ""

    decomposed = unicodedata.normalize("NFKD", value.strip().casefold())
    without_marks = "".join(
        char for char in decomposed if not unicodedata.combining(char)
    )

    return _HEADER_SEPARATORS.sub(" ", without_marks).strip()


def _map_header(header: list[Any]) -> tuple[dict[str, int], tuple[int, ...]]:
    if not header:
        raise ImportPreviewError("missing_header")

    mapping: dict[str, int] = {}
    ignored = []

    for index, cell in enumerate(header):
        if _is_blank_cell(cell):
            raise ImportPreviewError("blank_header")

        canonical = HEADER_ALIASES.get(_normalize_header(cell))

        if canonical is None:
            ignored.append(index)
            continue

        if canonical in mapping:
            raise ImportPreviewError("ambiguous_header")

        mapping[canonical] = index

    if not set(REQUIRED_FIELDS) <= set(mapping):
        raise ImportPreviewError("missing_required_fields")

    return mapping, tuple(ignored)


def _apply_manual_mapping(
    header: list[Any],
    manual: dict[str, int],
) -> tuple[dict[str, int], tuple[int, ...]]:
    """
    Confere o mapeamento manual contra o cabeçalho. Colunas selecionadas
    precisam existir e ter nome; as demais são ignoradas.
    """

    if not header:
        raise ImportPreviewError("missing_header")

    for index in manual.values():
        if index >= len(header) or _is_blank_cell(header[index]):
            raise ImportPreviewError("invalid_mapping")

    selected = set(manual.values())
    ignored = tuple(
        index for index in range(len(header)) if index not in selected
    )

    return dict(manual), ignored


# ---------------------------------------------------------------- linhas


def _validate_rows(
    records: list[list[Any] | None],
    mapping: dict[str, int],
    *,
    header_width: int,
) -> tuple[PreviewRow, ...]:
    rows = []
    seen: dict[tuple[str, Decimal, date], int] = {}

    for offset, record in enumerate(records):
        if record is None:
            continue

        row_number = offset + 2
        issues: list[RowIssue] = []

        if len(_trim_trailing_blanks(record)) > header_width:
            issues.append(RowIssue(None, "extra_columns"))

        def cell(field: str) -> Any:
            index = mapping.get(field)

            if index is None or index >= len(record):
                return None

            return record[index]

        values: dict[str, Any] = {}

        for field, converter in _CONVERTERS:
            try:
                values[field] = converter(cell(field))
            except _RowIssueError as issue:
                issues.append(RowIssue(field, issue.code))

        if not issues:
            key = (
                " ".join(values["cliente"].split()).casefold(),
                values["valor"],
                values["vencimento"],
            )

            if key in seen:
                issues.append(RowIssue(None, "duplicate_in_file"))
            else:
                seen[key] = row_number

        rows.append(
            PreviewRow(
                row_number=row_number,
                values=None if issues else values,
                issues=tuple(issues),
            )
        )

    return tuple(rows)


class _RowIssueError(Exception):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def _text_cell(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, bool):
        raise _RowIssueError("invalid_type")

    if isinstance(value, int):
        return str(value)

    if isinstance(value, str):
        return value.strip()

    raise _RowIssueError("invalid_type")


def _convert_cliente(value: Any) -> str:
    text = _text_cell(value)

    if not text:
        raise _RowIssueError("required")

    if _CONTROL_CHARS.search(text) or text.startswith(_FORMULA_PREFIXES):
        raise _RowIssueError("invalid_characters")

    if len(text) < CLIENTE_MIN_CHARS:
        raise _RowIssueError("too_short")

    if len(text) > CLIENTE_MAX_CHARS:
        raise _RowIssueError("too_long")

    return text


def _convert_email(value: Any) -> str | None:
    text = _text_cell(value)

    if not text:
        return None

    if len(text) > EMAIL_MAX_CHARS:
        raise _RowIssueError("too_long")

    if (
        _CONTROL_CHARS.search(text)
        or text.startswith(_FORMULA_PREFIXES)
        or not _EMAIL_PATTERN.match(text)
    ):
        raise _RowIssueError("invalid_email")

    return text


def _convert_whatsapp(value: Any) -> str | None:
    text = _text_cell(value)

    if not text:
        return None

    if len(text) > WHATSAPP_MAX_CHARS:
        raise _RowIssueError("too_long")

    digits = sum(1 for char in text if char.isdigit())

    if not _WHATSAPP_PATTERN.match(text) or digits < WHATSAPP_MIN_DIGITS:
        raise _RowIssueError("invalid_whatsapp")

    return text


def _convert_valor(value: Any) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise _RowIssueError("required")

    if isinstance(value, bool):
        raise _RowIssueError("invalid_type")

    if isinstance(value, (int, float, Decimal)):
        amount = _numeric_money(value)
    elif isinstance(value, str):
        amount = _text_money(value)
    else:
        raise _RowIssueError("invalid_type")

    if amount <= 0:
        raise _RowIssueError("non_positive_money")

    if amount > MAX_MONEY:
        raise _RowIssueError("money_out_of_range")

    return amount


def _numeric_money(value: int | float | Decimal) -> Decimal:
    try:
        amount = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
    except (InvalidOperation, ValueError) as error:
        raise _RowIssueError("invalid_money") from error

    if not amount.is_finite():
        raise _RowIssueError("invalid_money")

    if amount.as_tuple().exponent < -2:
        raise _RowIssueError("invalid_money")

    if abs(amount) > MAX_MONEY:
        raise _RowIssueError("money_out_of_range")

    return amount.quantize(Decimal("0.01"))


def _text_money(value: str) -> Decimal:
    text = value.strip()

    if text.startswith("R$"):
        text = text[2:]

    text = text.replace(" ", "").replace(" ", "")

    if text.startswith("-"):
        raise _RowIssueError("non_positive_money")

    if _MONEY_PLAIN.match(text):
        normalized = text.replace(",", ".")
    elif _MONEY_GROUPED_BR.match(text) or _MONEY_GROUPED_US.match(text):
        # Um único grupo de três dígitos sem parte decimal ("12.400" ou
        # "1,234") pode ser milhar ou decimal: rejeição conservadora.
        if sum(1 for char in text if char in ".,") == 1:
            raise _RowIssueError("ambiguous_money")

        normalized = text
    else:
        raise _RowIssueError("invalid_money")

    if Decimal(_strip_grouping(normalized)) > MAX_MONEY:
        raise _RowIssueError("money_out_of_range")

    try:
        return parse_localized_money(normalized)
    except ValueError as error:
        raise _RowIssueError("invalid_money") from error


def _strip_grouping(text: str) -> str:
    """
    Remove separadores de milhar de um texto já validado por regex,
    mantendo o último separador como ponto decimal quando houver.
    """

    head, separator, tail = max(
        (text.rpartition(","), text.rpartition(".")),
        key=lambda parts: len(parts[0]),
    )

    if not separator or len(tail) == 3:
        return text.replace(",", "").replace(".", "")

    return head.replace(",", "").replace(".", "") + "." + tail


def _convert_vencimento(value: Any) -> date:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise _RowIssueError("required")

    if isinstance(value, datetime):
        if value.time() != datetime.min.time() or value.tzinfo is not None:
            raise _RowIssueError("invalid_date")

        return value.date()

    if isinstance(value, date):
        return value

    if not isinstance(value, str):
        raise _RowIssueError("invalid_type")

    text = value.strip()

    if _ISO_DATE.match(text):
        pattern = "%Y-%m-%d"
    elif _BR_DATE.match(text):
        pattern = "%d/%m/%Y"
    else:
        raise _RowIssueError("invalid_date")

    try:
        return datetime.strptime(text, pattern).date()
    except ValueError as error:
        raise _RowIssueError("invalid_date") from error


def _convert_status(value: Any) -> str:
    if value is None:
        return "aberto"

    if not isinstance(value, str):
        raise _RowIssueError("invalid_type")

    text = value.strip().casefold()

    if text in ("", "aberto"):
        return "aberto"

    if text in ("pago", "atrasado"):
        raise _RowIssueError("status_not_allowed")

    raise _RowIssueError("invalid_status")


_CONVERTERS = (
    ("cliente", _convert_cliente),
    ("valor", _convert_valor),
    ("vencimento", _convert_vencimento),
    ("status", _convert_status),
    ("email", _convert_email),
    ("whatsapp", _convert_whatsapp),
)


# ---------------------------------------------------------------- digest


def _content_digest(
    mapping: dict[str, int],
    rows: tuple[PreviewRow, ...],
) -> str:
    """
    Digest determinístico do conteúdo normalizado, independente do
    formato (CSV/XLSX), delimitador, BOM e ordem das colunas.
    """

    payload = {
        "fields": sorted(mapping),
        "rows": [
            {
                "row": row.row_number,
                "values": None if row.values is None else {
                    field: (
                        None if row.values[field] is None
                        else str(row.values[field])
                    )
                    for field in CANONICAL_FIELDS
                },
                "issues": [
                    [issue.field, issue.code] for issue in row.issues
                ],
            }
            for row in rows
        ],
    }
    encoded = json.dumps(
        payload,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
    ).encode("ascii")

    return hashlib.sha256(encoded).hexdigest()
