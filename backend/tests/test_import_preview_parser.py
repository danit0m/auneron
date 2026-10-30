"""
Testes offline do parser de pré-visualização de importação CSV/XLSX.

Sem banco, rede ou disco: todos os arquivos são sintéticos e montados em
memória. O módulo é carregado diretamente pelo caminho do arquivo para não
executar ``app/services/__init__.py``, que importa serviços acoplados ao
banco.
"""

import ast
import builtins
import importlib.util
import io
import sys
import zipfile
from datetime import date
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest


BACKEND_DIR = Path(__file__).resolve().parents[1]
MODULE_PATH = (
    BACKEND_DIR / "app" / "services" / "import_preview_parser.py"
)

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def _load_parser_module():
    name = "import_preview_parser_under_test"
    spec = importlib.util.spec_from_file_location(name, MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


parser = _load_parser_module()
ImportPreviewError = parser.ImportPreviewError
parse = parser.parse_import_preview

PII_SENTINEL = "Fulana Sigilosa da Silva"
HEADER = "cliente;valor;vencimento;status;email;whatsapp"


def _csv(*lines, bom=False, encoding="utf-8"):
    data = "\r\n".join(lines).encode(encoding)
    return (b"\xef\xbb\xbf" + data) if bom else data


def _xlsx(rows, extra_sheet=False):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet("Outra")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rezip(data, *, add=None, replace=None, drop=()):
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename in drop:
                continue
            content = source.read(entry.filename)
            if replace and entry.filename in replace:
                content = replace[entry.filename](content)
            target.writestr(entry.filename, content)
        for name, content in (add or {}).items():
            target.writestr(name, content)
    return buffer.getvalue()


def _valid_xlsx():
    return _xlsx([
        ["cliente", "valor", "vencimento", "status"],
        ["Ana Souza", 1500, datetime(2026, 1, 5), "aberto"],
    ])


def _error_code(filename, data):
    with pytest.raises(ImportPreviewError) as caught:
        parse(filename, data)
    return caught.value.code


def _single_row_issues(row_line, header=HEADER):
    preview = parse("contas.csv", _csv(header, row_line))
    assert preview.total_rows == 1
    return [(issue.field, issue.code) for issue in preview.rows[0].issues]


# ------------------------------------------------------------ CSV válido


def test_csv_semicolon_with_bom_normalizes_row():
    preview = parse(
        "contas.csv",
        _csv(
            HEADER,
            "Ana Souza;R$ 1.234,56;05/01/2026;Aberto;ana@example.com;"
            "+55 (11) 99999-0000",
            bom=True,
        ),
    )

    assert preview.file_kind == "csv"
    assert preview.valid_rows == 1
    row = preview.rows[0]
    assert row.row_number == 2
    assert row.values == {
        "cliente": "Ana Souza",
        "valor": Decimal("1234.56"),
        "vencimento": date(2026, 1, 5),
        "status": "aberto",
        "email": "ana@example.com",
        "whatsapp": "+55 (11) 99999-0000",
    }
    assert preview.valid_total == Decimal("1234.56")


def test_csv_comma_delimiter_and_optional_fields_absent():
    preview = parse(
        "contas.csv",
        _csv("cliente,valor,vencimento,status", "Bruno Lima,250.00,2026-03-10,"),
    )

    row = preview.rows[0]
    assert row.is_valid
    assert row.values["email"] is None
    assert row.values["whatsapp"] is None
    assert row.values["status"] == "aberto"
    assert row.values["valor"] == Decimal("250.00")


def test_header_aliases_accents_and_case_are_mapped():
    preview = parse(
        "contas.csv",
        _csv(
            "Nome do Cliente;Valor;Data de Vencimento;Situação;E-mail;Observação",
            "Carla Dias;10,00;2026-02-01;aberto;c@example.com;qualquer",
        ),
    )

    assert preview.column_mapping == {
        "cliente": 0,
        "valor": 1,
        "vencimento": 2,
        "status": 3,
        "email": 4,
    }
    assert preview.ignored_columns == (5,)
    assert preview.rows[0].is_valid


def test_blank_lines_are_skipped_and_row_numbers_preserved():
    preview = parse(
        "contas.csv",
        _csv(
            HEADER,
            "Ana Souza;10,00;2026-01-01;aberto;;",
            ";;;;;",
            "",
            "Bruno Lima;20,00;2026-01-02;aberto;;",
            "",
        ),
    )

    assert [row.row_number for row in preview.rows] == [2, 5]
    assert preview.valid_total == Decimal("30.00")


# ------------------------------------------------------ CSV rejeitado


@pytest.mark.parametrize(
    "filename",
    ["contas.txt", "contas.xls", "contas.xlsm", "contas", "", None, "csv"],
)
def test_unknown_extensions_are_rejected(filename):
    assert _error_code(filename, b"cliente") == "unsupported_extension"


def test_empty_file_is_rejected():
    assert _error_code("contas.csv", b"") == "empty_file"


def test_file_too_large_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_FILE_BYTES", 10)
    assert _error_code("contas.csv", b"x" * 11) == "file_too_large"


def test_non_utf8_csv_is_rejected():
    data = _csv(HEADER, "José;10,00;2026-01-01;aberto;;", encoding="cp1252")
    assert _error_code("contas.csv", data) == "unsupported_encoding"


def test_nul_byte_is_rejected():
    assert _error_code("contas.csv", _csv(HEADER, "A\x00B;1;2026-01-01;;;")) == (
        "binary_content"
    )


def test_malformed_quotes_are_rejected():
    data = _csv(HEADER, '"Ana" Souza;10,00;2026-01-01;aberto;;')
    assert _error_code("contas.csv", data) == "malformed_csv"


def test_csv_extension_with_zip_content_is_rejected():
    assert _error_code("contas.csv", _valid_xlsx()) == "content_mismatch"


def test_header_only_is_rejected():
    assert _error_code("contas.csv", _csv(HEADER)) == "no_data_rows"


def test_missing_required_field_is_rejected():
    data = _csv("cliente;valor;vencimento", "Ana Souza;10,00;2026-01-01")
    assert _error_code("contas.csv", data) == "missing_required_fields"


def test_two_columns_for_same_field_are_ambiguous():
    data = _csv(
        "cliente;nome;valor;vencimento;status",
        "Ana Souza;Ana;10,00;2026-01-01;aberto",
    )
    assert _error_code("contas.csv", data) == "ambiguous_header"


def test_blank_header_cell_is_rejected():
    data = _csv(
        "cliente;;valor;vencimento;status",
        "Ana Souza;x;10,00;2026-01-01;aberto",
    )
    assert _error_code("contas.csv", data) == "blank_header"


def test_too_many_rows_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_DATA_ROWS", 2)
    lines = [HEADER] + [f"Cliente {i};1{i},00;2026-01-01;aberto;;" for i in range(3)]
    assert _error_code("contas.csv", _csv(*lines)) == "too_many_rows"


def test_too_many_columns_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_COLUMNS", 6)
    data = _csv(HEADER + ";extra", "Ana Souza;10,00;2026-01-01;aberto;;;x")
    assert _error_code("contas.csv", data) == "too_many_columns"


def test_cell_too_long_is_rejected():
    data = _csv(HEADER, "Ana Souza;10,00;2026-01-01;aberto;;" + "9" * 1001)
    assert _error_code("contas.csv", data) == "cell_too_long"


# --------------------------------------------------- validação por linha


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("12400", Decimal("12400.00")),
        ("12400,5", Decimal("12400.50")),
        ("12400.50", Decimal("12400.50")),
        ("12.400,00", Decimal("12400.00")),
        ("1.234.567", Decimal("1234567.00")),
        ("12,400.00", Decimal("12400.00")),
        ("R$ 1.234,56", Decimal("1234.56")),
        ("0,01", Decimal("0.01")),
    ],
)
def test_valid_money_formats(raw, expected):
    preview = parse("contas.csv", _csv(HEADER, f"Ana Souza;{raw};2026-01-01;;;"))
    assert preview.rows[0].values["valor"] == expected


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("", "required"),
        ("abc", "invalid_money"),
        ("10,555", "ambiguous_money"),
        ("12.400", "ambiguous_money"),
        ("1,234", "ambiguous_money"),
        ("1.234,567", "invalid_money"),
        ("1e5", "invalid_money"),
        ("+10", "invalid_money"),
        ("-10,00", "non_positive_money"),
        ("0,00", "non_positive_money"),
        ("0", "non_positive_money"),
        ("1.000.000.000.000,00", "money_out_of_range"),
        ("1000000000000", "money_out_of_range"),
        ("NaN", "invalid_money"),
        ("Infinity", "invalid_money"),
    ],
)
def test_invalid_money_is_row_error(raw, code):
    assert _single_row_issues(f"Ana Souza;{raw};2026-01-01;;;") == [("valor", code)]


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("2026-01-05", date(2026, 1, 5)), ("05/01/2026", date(2026, 1, 5))],
)
def test_valid_dates(raw, expected):
    preview = parse("contas.csv", _csv(HEADER, f"Ana Souza;10,00;{raw};;;"))
    assert preview.rows[0].values["vencimento"] == expected


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("", "required"),
        ("31/02/2026", "invalid_date"),
        ("2026/01/05", "invalid_date"),
        ("05-01-2026", "invalid_date"),
        ("5/1/2026", "invalid_date"),
        ("2026-1-5", "invalid_date"),
        ("05/01/26", "invalid_date"),
    ],
)
def test_invalid_dates_are_row_errors(raw, code):
    assert _single_row_issues(f"Ana Souza;10,00;{raw};;;") == [("vencimento", code)]


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("pago", "status_not_allowed"),
        ("Atrasado", "status_not_allowed"),
        ("cancelado", "invalid_status"),
    ],
)
def test_initial_status_other_than_aberto_is_rejected(raw, code):
    assert _single_row_issues(f"Ana Souza;10,00;2026-01-01;{raw};;") == [
        ("status", code)
    ]


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("", "required"),
        ("A", "too_short"),
        ("A" * 151, "too_long"),
        ("=HYPERLINK(1)", "invalid_characters"),
        ("@SUM(1)", "invalid_characters"),
        ('"Ana\nSouza"', "invalid_characters"),
    ],
)
def test_invalid_cliente_is_row_error(raw, code):
    assert _single_row_issues(f"{raw};10,00;2026-01-01;;;") == [("cliente", code)]


def test_cliente_at_max_length_is_valid():
    preview = parse("contas.csv", _csv(HEADER, "A" * 150 + ";10,00;2026-01-01;;;"))
    assert preview.rows[0].is_valid


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("sem-arroba", "invalid_email"),
        ("a@b", "invalid_email"),
        ("=a@b.com", "invalid_email"),
        ("a" * 140 + "@example.com", "too_long"),
    ],
)
def test_invalid_email_is_row_error(raw, code):
    assert _single_row_issues(f"Ana Souza;10,00;2026-01-01;;{raw};") == [
        ("email", code)
    ]


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("-11999990000", "invalid_whatsapp"),
        ("=1+1", "invalid_whatsapp"),
        ("1234567", "invalid_whatsapp"),
        ("(--------)", "invalid_whatsapp"),
        ("1" * 31, "too_long"),
    ],
)
def test_invalid_whatsapp_is_row_error(raw, code):
    assert _single_row_issues(f"Ana Souza;10,00;2026-01-01;;;{raw}") == [
        ("whatsapp", code)
    ]


def test_row_with_more_columns_than_header_is_flagged():
    assert _single_row_issues(
        "Ana Souza;10,00;2026-01-01;aberto",
        header="cliente;valor;vencimento;status",
    ) == []
    assert _single_row_issues(
        "Ana Souza;10,00;2026-01-01;aberto;sobra",
        header="cliente;valor;vencimento;status",
    ) == [(None, "extra_columns")]


def test_short_row_reports_required_fields():
    issues = _single_row_issues("Ana Souza")
    assert ("valor", "required") in issues
    assert ("vencimento", "required") in issues


def test_multiple_issues_are_all_reported_and_values_withheld():
    preview = parse("contas.csv", _csv(HEADER, "A;abc;99/99/2026;pago;x;y"))
    row = preview.rows[0]
    assert not row.is_valid
    assert row.values is None
    assert {issue.field for issue in row.issues} == {
        "cliente", "valor", "vencimento", "status", "email", "whatsapp",
    }


def test_duplicate_in_file_flags_later_occurrence_only():
    preview = parse(
        "contas.csv",
        _csv(
            HEADER,
            "Ana Souza;10,00;2026-01-01;aberto;;",
            "Bruno Lima;10,00;2026-01-01;aberto;;",
            "  ana   SOUZA ;10.00;01/01/2026;;;",
        ),
    )

    assert [row.is_valid for row in preview.rows] == [True, True, False]
    assert [(i.field, i.code) for i in preview.rows[2].issues] == [
        (None, "duplicate_in_file")
    ]
    assert preview.valid_rows == 2
    assert preview.invalid_rows == 1


# ------------------------------------------------------- sem dados pessoais


def test_issues_and_file_errors_do_not_echo_raw_content():
    preview = parse(
        "contas.csv",
        _csv(HEADER, f"{PII_SENTINEL};{PII_SENTINEL};{PII_SENTINEL};"
             f"{PII_SENTINEL};{PII_SENTINEL};{PII_SENTINEL}"),
    )
    rendered = repr(preview.rows[0].issues) + " ".join(
        issue.message for issue in preview.rows[0].issues
    )
    assert PII_SENTINEL not in rendered
    assert PII_SENTINEL.split()[0] not in rendered

    with pytest.raises(ImportPreviewError) as caught:
        parse(f"{PII_SENTINEL}.csv", _csv(f"{PII_SENTINEL};valor", "x;1"))
    error = caught.value
    for text in (str(error), repr(error), error.message, repr(error.args)):
        assert PII_SENTINEL not in text


# ------------------------------------------------------------- digest


def test_digests_are_deterministic_and_format_independent():
    rows = ["Ana Souza;10.00;2026-01-01;aberto;;", "Bruno Lima;20.00;2026-01-02;;;"]
    semicolon = _csv(HEADER, *rows)
    comma = _csv(*(line.replace(";", ",") for line in [HEADER, *rows]))
    with_bom = _csv(HEADER, *rows, bom=True)

    first = parse("contas.csv", semicolon)
    again = parse("contas.csv", semicolon)
    other_delimiter = parse("contas.csv", comma)
    bom = parse("contas.csv", with_bom)

    assert first.file_sha256 == again.file_sha256
    assert first.content_sha256 == again.content_sha256
    assert len(first.content_sha256) == 64
    assert other_delimiter.file_sha256 != first.file_sha256
    assert other_delimiter.content_sha256 == first.content_sha256
    assert bom.content_sha256 == first.content_sha256

    xlsx = parse("contas.xlsx", _xlsx([
        ["cliente", "valor", "vencimento", "status", "email", "whatsapp"],
        ["Ana Souza", 10, datetime(2026, 1, 1), "aberto", None, None],
        ["Bruno Lima", 20, datetime(2026, 1, 2), None, None, None],
    ]))
    assert xlsx.content_sha256 == first.content_sha256


def test_content_digest_changes_with_values():
    base = parse("contas.csv", _csv(HEADER, "Ana Souza;10,00;2026-01-01;;;"))
    changed = parse("contas.csv", _csv(HEADER, "Ana Souza;10,01;2026-01-01;;;"))
    assert base.content_sha256 != changed.content_sha256


# ---------------------------------------------------------------- XLSX


def test_valid_xlsx_with_typed_cells():
    preview = parse("Contas.XLSX", _xlsx([
        ["Cliente", "Valor", "Vencimento", "Status", "WhatsApp"],
        ["Ana Souza", 1500, datetime(2026, 1, 5), "aberto", 5511999990000],
        ["Bruno Lima", 99.9, "05/02/2026", None, None],
        ["Carla Dias", "1.234,50", "2026-03-05", "Aberto", None],
    ]))

    assert preview.file_kind == "xlsx"
    assert preview.valid_rows == 3
    assert preview.rows[0].values["valor"] == Decimal("1500.00")
    assert preview.rows[0].values["vencimento"] == date(2026, 1, 5)
    assert preview.rows[0].values["whatsapp"] == "5511999990000"
    assert preview.rows[1].values["valor"] == Decimal("99.90")
    assert preview.rows[2].values["valor"] == Decimal("1234.50")
    assert preview.valid_total == Decimal("2834.40")


def test_xlsx_typed_cell_rejections():
    preview = parse("contas.xlsx", _xlsx([
        ["cliente", "valor", "vencimento", "status"],
        ["Ana Souza", 10.555, datetime(2026, 1, 5), "aberto"],
        ["Bruno Lima", True, datetime(2026, 1, 5), "aberto"],
        ["Carla Dias", 10, datetime(2026, 1, 5, 10, 30), "aberto"],
        ["Davi Melo", 10, 46000, "aberto"],
        [3.5, 10, datetime(2026, 1, 5), "aberto"],
    ]))

    assert [
        [(i.field, i.code) for i in row.issues] for row in preview.rows
    ] == [
        [("valor", "invalid_money")],
        [("valor", "invalid_type")],
        [("vencimento", "invalid_date")],
        [("vencimento", "invalid_type")],
        [("cliente", "invalid_type")],
    ]


def test_xlsx_header_only_is_rejected():
    data = _xlsx([["cliente", "valor", "vencimento", "status"]])
    assert _error_code("contas.xlsx", data) == "no_data_rows"


def test_xlsx_formula_is_rejected():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["cliente", "valor", "vencimento", "status"])
    sheet.append(["Ana Souza", "=1+1", "2026-01-01", "aberto"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert _error_code("contas.xlsx", buffer.getvalue()) == "formula_workbook"


def test_xlsx_multiple_sheets_are_rejected():
    data = _xlsx(
        [["cliente", "valor", "vencimento", "status"], ["Ana", 1, "2026-01-01", ""]],
        extra_sheet=True,
    )
    assert _error_code("contas.xlsx", data) == "too_many_sheets"


def test_xlsx_with_vba_project_is_rejected():
    data = _rezip(_valid_xlsx(), add={"xl/vbaProject.bin": b"\x00" * 16})
    assert _error_code("contas.xlsx", data) == "macro_workbook"


def test_xlsx_with_macro_enabled_content_type_is_rejected():
    data = _rezip(
        _valid_xlsx(),
        replace={
            "[Content_Types].xml": lambda content: content.replace(
                b"sheet.main+xml", b"sheet.macroEnabled.main+xml"
            )
        },
    )
    assert _error_code("contas.xlsx", data) == "macro_workbook"


def test_xlsx_with_external_links_is_rejected():
    data = _rezip(
        _valid_xlsx(),
        add={"xl/externalLinks/externalLink1.xml": b"<externalLink/>"},
    )
    assert _error_code("contas.xlsx", data) == "external_links"


def test_xlsx_with_doctype_is_rejected():
    data = _rezip(
        _valid_xlsx(),
        add={"docProps/custom.xml": b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><x/>'},
    )
    assert _error_code("contas.xlsx", data) == "unsafe_xml"


def test_xlsx_without_workbook_part_is_rejected():
    data = _rezip(_valid_xlsx(), drop=("xl/workbook.xml",))
    assert _error_code("contas.xlsx", data) == "malformed_workbook"


@pytest.mark.parametrize(
    "data",
    [b"PK\x03\x04" + b"\x00" * 64, b"PK\x03\x04not-a-zip-at-all"],
)
def test_malformed_zip_is_rejected(data):
    assert _error_code("contas.xlsx", data) == "malformed_zip"


def test_xlsx_extension_with_text_content_is_rejected():
    assert _error_code("contas.xlsx", _csv(HEADER, "Ana;1;2026-01-01;;;")) == (
        "content_mismatch"
    )


def test_xlsx_compression_bomb_is_rejected():
    data = _rezip(_valid_xlsx(), add={"xl/media/pad.bin": b"0" * (3 * 1024 * 1024)})
    assert _error_code("contas.xlsx", data) == "unsafe_zip"


def test_xlsx_too_many_zip_entries_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_ZIP_ENTRIES", 3)
    assert _error_code("contas.xlsx", _valid_xlsx()) == "unsafe_zip"


def test_xlsx_encrypted_entry_is_rejected():
    data = bytearray(_valid_xlsx())
    central = data.index(b"PK\x01\x02")
    data[central + 8] |= 0x01
    assert _error_code("contas.xlsx", bytes(data)) == "encrypted_workbook"


def test_xlsx_too_many_rows_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_DATA_ROWS", 2)
    rows = [["cliente", "valor", "vencimento", "status"]] + [
        [f"Cliente {i}", 10 + i, "2026-01-01", "aberto"] for i in range(3)
    ]
    assert _error_code("contas.xlsx", _xlsx(rows)) == "too_many_rows"


# ---------------------------------------------- mapeamento manual de colunas


ARBITRARY_HEADER = "Devedor;Montante;Prazo;Obs;Situacao atual"
ARBITRARY_MAPPING = {"cliente": 0, "valor": 1, "vencimento": 2, "status": 4}


def _manual_error_code(data, mapping, delimiter=";", filename="contas.csv"):
    with pytest.raises(ImportPreviewError) as caught:
        parse(
            filename,
            data,
            column_mapping=mapping,
            csv_delimiter=delimiter,
        )
    return caught.value.code


def _arbitrary_csv():
    return _csv(ARBITRARY_HEADER, "Ana Souza;10,00;2026-01-01;qualquer;")


def test_manual_mapping_csv_semicolon_with_arbitrary_headers():
    preview = parse(
        "contas.csv",
        _csv(
            ARBITRARY_HEADER,
            "Ana Souza;1.234,56;05/01/2026;qualquer;",
            "Bruno Lima;20,00;2026-01-02;;Aberto",
        ),
        column_mapping=ARBITRARY_MAPPING,
        csv_delimiter=";",
    )

    assert preview.column_mapping == ARBITRARY_MAPPING
    assert preview.ignored_columns == (3,)
    assert preview.valid_rows == 2
    assert preview.rows[0].values == {
        "cliente": "Ana Souza",
        "valor": Decimal("1234.56"),
        "vencimento": date(2026, 1, 5),
        "status": "aberto",
        "email": None,
        "whatsapp": None,
    }
    assert preview.valid_total == Decimal("1254.56")


def test_manual_mapping_csv_comma_with_unknown_names_and_quoted_decimal():
    preview = parse(
        "contas.csv",
        _csv(
            "Nome Completo,Quantia,Prazo,Estado,Contato",
            'Ana Souza,"1.234,56",2026-01-05,aberto,ana@example.com',
        ),
        column_mapping={
            "cliente": 0, "valor": 1, "vencimento": 2, "status": 3, "email": 4,
        },
        csv_delimiter=",",
    )

    assert preview.ignored_columns == ()
    row = preview.rows[0]
    assert row.is_valid
    assert row.values["valor"] == Decimal("1234.56")
    assert row.values["email"] == "ana@example.com"


def test_manual_mapping_xlsx_with_arbitrary_headers():
    preview = parse(
        "contas.xlsx",
        _xlsx([
            ["Devedor", "Montante", "Prazo", "Estado", "Telefone", "Notas"],
            ["Ana Souza", 1500, datetime(2026, 1, 5), "aberto", 5511999990000, "x"],
            ["Bruno Lima", "99,90", "05/02/2026", None, None, None],
        ]),
        column_mapping={
            "cliente": 0, "valor": 1, "vencimento": 2, "status": 3, "whatsapp": 4,
        },
    )

    assert preview.file_kind == "xlsx"
    assert preview.ignored_columns == (5,)
    assert preview.valid_rows == 2
    assert preview.rows[0].values["whatsapp"] == "5511999990000"
    assert preview.rows[1].values["status"] == "aberto"
    assert preview.valid_total == Decimal("1599.90")


def test_manual_mapping_with_reordered_columns_keeps_content_digest():
    auto = parse(
        "contas.csv",
        _csv(
            "cliente;valor;vencimento;status",
            "Ana Souza;10.00;2026-01-01;aberto",
            "Bruno Lima;20.00;2026-01-02;",
        ),
    )
    manual_mapping = {"cliente": 4, "valor": 2, "vencimento": 0, "status": 3}
    manual = parse(
        "contas.csv",
        _csv(
            "Prazo,Obs,Montante,Estado,Devedor",
            "2026-01-01,a,10.00,aberto,Ana Souza",
            "2026-01-02,b,20.00,,Bruno Lima",
        ),
        column_mapping=manual_mapping,
        csv_delimiter=",",
    )
    xlsx = parse(
        "contas.xlsx",
        _xlsx([
            ["Estado", "Devedor", "Prazo", "Montante"],
            ["aberto", "Ana Souza", datetime(2026, 1, 1), 10],
            [None, "Bruno Lima", datetime(2026, 1, 2), 20],
        ]),
        column_mapping={"cliente": 1, "valor": 3, "vencimento": 2, "status": 0},
    )

    assert manual.column_mapping == manual_mapping
    assert manual.ignored_columns == (1,)
    assert manual.content_sha256 == auto.content_sha256
    assert xlsx.content_sha256 == auto.content_sha256
    assert manual.file_sha256 != auto.file_sha256


def test_manual_mapping_detects_duplicates():
    preview = parse(
        "contas.csv",
        _csv(
            ARBITRARY_HEADER,
            "Ana Souza;10,00;2026-01-01;;",
            "  ana   SOUZA ;10.00;01/01/2026;outra;",
        ),
        column_mapping=ARBITRARY_MAPPING,
        csv_delimiter=";",
    )

    assert [row.is_valid for row in preview.rows] == [True, False]
    assert [(i.field, i.code) for i in preview.rows[1].issues] == [
        (None, "duplicate_in_file")
    ]


def test_manual_mapping_row_issues_still_apply():
    preview = parse(
        "contas.csv",
        _csv(ARBITRARY_HEADER, "Ana Souza;12.400;2026-01-01;;pago"),
        column_mapping=ARBITRARY_MAPPING,
        csv_delimiter=";",
    )

    assert [(i.field, i.code) for i in preview.rows[0].issues] == [
        ("valor", "ambiguous_money"),
        ("status", "status_not_allowed"),
    ]


def test_without_mapping_arbitrary_headers_keep_default_error():
    assert _error_code("contas.csv", _arbitrary_csv()) == "missing_required_fields"


def test_manual_mapping_header_only_is_rejected():
    assert _manual_error_code(_csv(ARBITRARY_HEADER), ARBITRARY_MAPPING) == (
        "no_data_rows"
    )


def test_manual_mapping_with_malformed_csv_is_rejected():
    data = _csv(ARBITRARY_HEADER, '"Ana" Souza;10,00;2026-01-01;;')
    assert _manual_error_code(data, ARBITRARY_MAPPING) == "malformed_csv"


@pytest.mark.parametrize(
    "mapping",
    [
        {},
        {"cliente": 0, "valor": 1, "vencimento": 2},
        {"cliente": 0, "valor": 1, "vencimento": 2, "status": 4, "observacao": 3},
        {"Cliente": 0, "valor": 1, "vencimento": 2, "status": 4},
        {"cliente": True, "valor": 1, "vencimento": 2, "status": 4},
        {"cliente": 0, "valor": False, "vencimento": 2, "status": 4},
        {"cliente": 0.0, "valor": 1, "vencimento": 2, "status": 4},
        {"cliente": "0", "valor": 1, "vencimento": 2, "status": 4},
        {"cliente": -1, "valor": 1, "vencimento": 2, "status": 4},
        {"cliente": None, "valor": 1, "vencimento": 2, "status": 4},
        [("cliente", 0), ("valor", 1), ("vencimento", 2), ("status", 4)],
        "cliente",
    ],
)
def test_invalid_manual_mapping_is_rejected(mapping):
    assert _manual_error_code(_arbitrary_csv(), mapping) == "invalid_mapping"


def test_manual_mapping_with_repeated_indices_is_rejected():
    mapping = {"cliente": 0, "valor": 1, "vencimento": 2, "status": 1}
    assert _manual_error_code(_arbitrary_csv(), mapping) == "invalid_mapping"


@pytest.mark.parametrize("index", [5, 6, 30, 10**9])
def test_manual_mapping_index_outside_header_is_rejected(index):
    mapping = {**ARBITRARY_MAPPING, "email": index}
    assert _manual_error_code(_arbitrary_csv(), mapping) == "invalid_mapping"


def test_manual_mapping_index_on_trailing_blank_header_is_rejected():
    data = _csv(ARBITRARY_HEADER + ";", "Ana Souza;10,00;2026-01-01;;;x")
    mapping = {**ARBITRARY_MAPPING, "email": 5}
    assert _manual_error_code(data, mapping) == "invalid_mapping"


def test_manual_mapping_selecting_blank_header_is_rejected():
    data = _csv("Devedor;;Prazo;Situacao", "Ana Souza;10,00;2026-01-01;")
    mapping = {"cliente": 0, "valor": 1, "vencimento": 2, "status": 3}
    assert _manual_error_code(data, mapping) == "invalid_mapping"


def test_manual_mapping_ignores_unselected_blank_header():
    preview = parse(
        "contas.csv",
        _csv("Devedor;Montante;;Prazo;Situacao", "Ana Souza;10,00;x;2026-01-01;"),
        column_mapping={"cliente": 0, "valor": 1, "vencimento": 3, "status": 4},
        csv_delimiter=";",
    )

    assert preview.ignored_columns == (2,)
    assert preview.rows[0].is_valid


def test_manual_mapping_xlsx_selecting_blank_header_is_rejected():
    data = _xlsx([
        ["Devedor", None, "Prazo", "Estado"],
        ["Ana Souza", 10, datetime(2026, 1, 1), "aberto"],
    ])
    mapping = {"cliente": 0, "valor": 1, "vencimento": 2, "status": 3}
    assert _manual_error_code(data, mapping, None, "contas.xlsx") == (
        "invalid_mapping"
    )


@pytest.mark.parametrize("delimiter", [None, "", "\t", "|", ";;", " ;", 59, b";"])
def test_manual_mapping_csv_requires_valid_delimiter(delimiter):
    assert _manual_error_code(_arbitrary_csv(), ARBITRARY_MAPPING, delimiter) == (
        "invalid_delimiter"
    )


def test_manual_mapping_parses_only_with_given_delimiter():
    # Com "," o cabeçalho separado por ";" vira uma única coluna.
    assert _manual_error_code(_arbitrary_csv(), ARBITRARY_MAPPING, ",") == (
        "invalid_mapping"
    )


@pytest.mark.parametrize("delimiter", [";", ",", ""])
def test_delimiter_without_mapping_is_rejected(delimiter):
    data = _csv(HEADER, "Ana Souza;10,00;2026-01-01;;;")
    with pytest.raises(ImportPreviewError) as caught:
        parse("contas.csv", data, csv_delimiter=delimiter)
    assert caught.value.code == "invalid_delimiter"

    with pytest.raises(ImportPreviewError) as caught:
        parse("contas.xlsx", _valid_xlsx(), csv_delimiter=delimiter)
    assert caught.value.code == "invalid_delimiter"


@pytest.mark.parametrize("delimiter", [";", ",", ""])
def test_xlsx_with_csv_delimiter_is_rejected(delimiter):
    mapping = {"cliente": 0, "valor": 1, "vencimento": 2, "status": 3}
    assert _manual_error_code(_valid_xlsx(), mapping, delimiter, "contas.xlsx") == (
        "invalid_delimiter"
    )


def test_manual_options_are_keyword_only():
    with pytest.raises(TypeError):
        parse("contas.csv", _arbitrary_csv(), ARBITRARY_MAPPING, ";")


def test_manual_mapping_result_is_a_canonical_copy():
    mapping = {"status": 4, "vencimento": 2, "valor": 1, "cliente": 0}
    preview = parse(
        "contas.csv",
        _arbitrary_csv(),
        column_mapping=mapping,
        csv_delimiter=";",
    )
    mapping["cliente"] = 3

    assert list(preview.column_mapping) == [
        "cliente", "valor", "vencimento", "status",
    ]
    assert preview.column_mapping["cliente"] == 0
    assert preview.column_mapping is not mapping


def test_manual_mapping_errors_do_not_echo_raw_content():
    data = _csv(
        f"{PII_SENTINEL};{PII_SENTINEL};Prazo;Situacao",
        f"{PII_SENTINEL};{PII_SENTINEL};{PII_SENTINEL};{PII_SENTINEL}",
    )
    attempts = [
        ({PII_SENTINEL: 0, "valor": 1, "vencimento": 2, "status": 3}, ";"),
        ({"cliente": PII_SENTINEL, "valor": 1, "vencimento": 2, "status": 3}, ";"),
        ({"cliente": 0, "valor": 1, "vencimento": 2, "status": 9}, ";"),
        ({"cliente": 0, "valor": 1, "vencimento": 2, "status": 3}, PII_SENTINEL),
    ]

    for mapping, delimiter in attempts:
        with pytest.raises(ImportPreviewError) as caught:
            parse(
                f"{PII_SENTINEL}.csv",
                data,
                column_mapping=mapping,
                csv_delimiter=delimiter,
            )
        error = caught.value
        assert error.code in ("invalid_mapping", "invalid_delimiter")
        for text in (str(error), repr(error), error.message, repr(error.args)):
            assert PII_SENTINEL not in text
            assert PII_SENTINEL.split()[0] not in text

    preview = parse(
        "contas.csv",
        data,
        column_mapping={"cliente": 0, "valor": 1, "vencimento": 2, "status": 3},
        csv_delimiter=";",
    )
    rendered = repr(preview.rows[0].issues) + " ".join(
        issue.message for issue in preview.rows[0].issues
    )
    assert preview.rows[0].issues
    assert PII_SENTINEL not in rendered


# ------------------------------------------------- pureza / sem efeitos


ALLOWED_IMPORTS = {
    "__future__",
    "csv",
    "dataclasses",
    "datetime",
    "decimal",
    "hashlib",
    "io",
    "json",
    "openpyxl",
    "re",
    "typing",
    "unicodedata",
    "zipfile",
    "app.core.money",
}

FORBIDDEN_CALLS = {"open", "print", "eval", "exec", "compile", "__import__"}


def _module_tree():
    return ast.parse(MODULE_PATH.read_text(encoding="utf-8"))


def test_module_imports_only_pure_dependencies():
    imported = set()
    for node in ast.walk(_module_tree()):
        if isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom):
            imported.add(node.module)

    assert imported <= ALLOWED_IMPORTS, imported - ALLOWED_IMPORTS


def test_module_has_no_io_or_dynamic_calls():
    for node in ast.walk(_module_tree()):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            assert node.func.id not in FORBIDDEN_CALLS, node.func.id


def test_parsing_never_opens_files(monkeypatch):
    def _forbidden_open(*args, **kwargs):
        raise AssertionError("parser must not open files")

    csv_data = _csv(HEADER, "Ana Souza;10,00;2026-01-01;;;")
    xlsx_data = _valid_xlsx()

    monkeypatch.setattr(builtins, "open", _forbidden_open)
    monkeypatch.setattr(io, "open", _forbidden_open)

    assert parse("contas.csv", csv_data).valid_rows == 1
    assert parse("contas.xlsx", xlsx_data).valid_rows == 1


def test_non_bytes_input_is_rejected():
    with pytest.raises(TypeError):
        parse("contas.csv", "cliente;valor")
